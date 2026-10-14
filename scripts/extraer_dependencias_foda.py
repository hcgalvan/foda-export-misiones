#!/usr/bin/env python3
"""
extraer_dependencias_foda.py — FODA Export Misiones (Fase 3)

Extrae la tabla de dependencias internas->externas de la hoja `VariablesFoda`
del Excel original. Esta tabla es ESTRUCTURAL (como la tabla `Rango`): define,
para cada factor interno, de qué factores externos "depende" (con qué
factores externos se cruza en la matriz FODA). No depende de las respuestas
de una evaluación concreta -- es fija para el instrumento.

Corrección importante (ver ARCHITECTURE.md §5): se había documentado que el
cuadrante FODA dominante (`PART2 FODA`/`MATRIZ FODA`) requería carga
cualitativa manual por evaluación. Eso era INCORRECTO -- `PART2 FODA` es
100% fórmulas: cruza automáticamente cada factor interno con los factores
externos de los que depende (según esta tabla fija `VariablesFoda`) y
clasifica el cruce en FO/FA/DO/DA según las respuestas de situación de
ambos factores. Esta tabla de dependencias es lo único que hacía falta
extraer para poder reproducir ese cálculo en el motor.

Un factor interno con lista vacía es "Independiente": no participa en
ningún cruce (no aparece en PART2 FODA en absoluto).

Uso:
    python3 scripts/extraer_dependencias_foda.py \
        --excel docs/referencia/FODA_Export_Misiones_2026_DEMO.xlsx \
        --out frontend/src/data/dependenciasFoda.json
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instalá con:\n  pip install openpyxl --break-system-packages")

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_EXCEL = REPO_ROOT / "docs" / "referencia" / "FODA_Export_Misiones_2026_DEMO.xlsx"
DEFAULT_OUT = REPO_ROOT / "frontend" / "src" / "data" / "dependenciasFoda.json"


def extraer(excel_path: Path) -> dict:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["VariablesFoda"]

    dependencias = {}
    for r in range(3, ws.max_row + 1):
        codigo = ws.cell(row=r, column=3).value  # C: código del factor interno (p.ej. IA1)
        if not codigo:
            continue
        deps = []
        for col in range(6, 11):  # F..J: hasta 5 columnas de dependencia
            v = ws.cell(row=r, column=col).value
            if v:
                deps.append(str(v).strip())
        dependencias[str(codigo).strip()] = deps

    return dependencias


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    if not args.excel.exists():
        sys.exit(f"No se encontró el Excel en {args.excel}")

    dependencias = extraer(args.excel)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", encoding="utf-8") as f:
        json.dump(dependencias, f, ensure_ascii=False, indent=2)

    independientes = [k for k, v in dependencias.items() if not v]
    total_deps = sum(len(v) for v in dependencias.values())
    print(f"OK: {args.out} generado desde {args.excel.name}.")
    print(f"  Factores internos con tabla de dependencias: {len(dependencias)} (esperado: 64)")
    print(f"  Independientes (sin cruce): {len(independientes)} -> {independientes}")
    print(f"  Total de pares interno-externo (cruces posibles): {total_deps}")


if __name__ == "__main__":
    main()
